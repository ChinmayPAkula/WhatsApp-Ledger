"""
Excel ledger report generation — a "Ledger" table of entries in a date range,
plus a statistics block (total spend, spend by category, quantity by category,
price change by item/vendor) offset to the right on the same sheet.
"""

from datetime import date, datetime
from io import BytesIO
from collections import defaultdict

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from app.database import (
    get_entries_between,
    get_entries_before,
    get_all_stock_transactions,
    get_stock_transactions_between,
)


def resolve_period(period: str | None, reference: date = None) -> tuple[date, date]:
    """Turn a period token ('this_month' | 'last_month' | 'YYYY-MM' | None) into a [start, end) range."""
    today = reference or date.today()

    def month_bounds(year: int, month: int) -> tuple[date, date]:
        start = date(year, month, 1)
        end = date(year + 1, 1, 1) if month == 12 else date(year, month + 1, 1)
        return start, end

    if not period or period == "this_month":
        return month_bounds(today.year, today.month)
    if period == "last_month":
        y, m = (today.year - 1, 12) if today.month == 1 else (today.year, today.month - 1)
        return month_bounds(y, m)
    try:
        parsed = datetime.strptime(period, "%Y-%m")
        return month_bounds(parsed.year, parsed.month)
    except ValueError:
        return month_bounds(today.year, today.month)

MONEY_FORMAT = "₹#,##0.00"
BOLD = Font(bold=True)

LEDGER_HEADERS = [
    "Date", "Sender", "Vendor", "Item", "Category", "Quantity", "Unit",
    "Price/Unit", "Total Price", "Type", "Status", "Notes",
]
STATS_START_COL = 15  # column O — leaves a gap after the 12 ledger columns


def _entry_spend(entry: dict) -> float:
    if entry.get("total_price") is not None:
        return float(entry["total_price"])
    if entry.get("quantity") is not None and entry.get("price_per_unit") is not None:
        return float(entry["quantity"]) * float(entry["price_per_unit"])
    return 0.0


def _last_price_before(history: list[dict]) -> dict:
    """Map (item, vendor) -> most recent price_per_unit seen before the period."""
    last = {}
    for row in history:
        key = ((row.get("item") or "").strip().lower(), (row.get("vendor") or "").strip().lower())
        last[key] = row["price_per_unit"]  # history is oldest-first, so later rows overwrite
    return last


async def generate_report_workbook(start: date, end: date) -> bytes:
    entries = await get_entries_between(start, end)
    history = await get_entries_before(start)
    last_price = _last_price_before(history)

    wb = Workbook()
    ws = wb.active
    ws.title = "Ledger"

    # ── Ledger table ──
    for col, header in enumerate(LEDGER_HEADERS, start=1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = BOLD

    category_spend = defaultdict(float)
    category_qty = defaultdict(float)  # keyed by (category, unit)
    latest_price_in_period = {}  # (item, vendor) -> price_per_unit, in-period latest wins
    total_spend = 0.0

    for r, e in enumerate(entries, start=2):
        spend = _entry_spend(e)
        total_spend += spend
        category = e.get("category") or "Other"
        category_spend[category] += spend
        if e.get("quantity") is not None:
            category_qty[(category, e.get("unit") or "-")] += float(e["quantity"])

        item = (e.get("item") or "").strip().lower()
        vendor = (e.get("vendor") or "").strip().lower()
        if item and e.get("price_per_unit") is not None:
            latest_price_in_period[(item, vendor)] = e["price_per_unit"]

        ws.cell(row=r, column=1, value=e.get("created_at"))
        ws.cell(row=r, column=2, value=e.get("sender_phone"))
        ws.cell(row=r, column=3, value=e.get("vendor"))
        ws.cell(row=r, column=4, value=e.get("item"))
        ws.cell(row=r, column=5, value=e.get("category"))
        ws.cell(row=r, column=6, value=e.get("quantity"))
        ws.cell(row=r, column=7, value=e.get("unit"))
        c8 = ws.cell(row=r, column=8, value=e.get("price_per_unit"))
        c8.number_format = MONEY_FORMAT
        c9 = ws.cell(row=r, column=9, value=e.get("total_price"))
        c9.number_format = MONEY_FORMAT
        ws.cell(row=r, column=10, value=e.get("entry_type"))
        ws.cell(row=r, column=11, value=e.get("status"))
        ws.cell(row=r, column=12, value=e.get("notes"))

    for col in range(1, len(LEDGER_HEADERS) + 1):
        ws.column_dimensions[get_column_letter(col)].width = 16

    # ── Statistics block (offset to the right) ──
    sc = STATS_START_COL
    row = 1
    ws.cell(row=row, column=sc, value="Total Spend").font = BOLD
    cell = ws.cell(row=row, column=sc + 1, value=total_spend)
    cell.number_format = MONEY_FORMAT
    row += 2

    ws.cell(row=row, column=sc, value="Spend by Category").font = BOLD
    row += 1
    ws.cell(row=row, column=sc, value="Category").font = BOLD
    ws.cell(row=row, column=sc + 1, value="Spend").font = BOLD
    row += 1
    for cat, amount in sorted(category_spend.items(), key=lambda kv: kv[1], reverse=True):
        ws.cell(row=row, column=sc, value=cat)
        c = ws.cell(row=row, column=sc + 1, value=amount)
        c.number_format = MONEY_FORMAT
        row += 1
    row += 1

    ws.cell(row=row, column=sc, value="Quantity by Category").font = BOLD
    row += 1
    ws.cell(row=row, column=sc, value="Category").font = BOLD
    ws.cell(row=row, column=sc + 1, value="Unit").font = BOLD
    ws.cell(row=row, column=sc + 2, value="Total Qty").font = BOLD
    row += 1
    for (cat, unit), qty in sorted(category_qty.items(), key=lambda kv: kv[0]):
        ws.cell(row=row, column=sc, value=cat)
        ws.cell(row=row, column=sc + 1, value=unit)
        ws.cell(row=row, column=sc + 2, value=qty)
        row += 1
    row += 1

    ws.cell(row=row, column=sc, value="Price Change (vs. last known price)").font = BOLD
    row += 1
    headers = ["Item", "Vendor", "Previous Price", "Current Price", "Change", "% Change", "Direction"]
    for i, h in enumerate(headers):
        ws.cell(row=row, column=sc + i, value=h).font = BOLD
    row += 1
    for (item, vendor), current in sorted(latest_price_in_period.items()):
        previous = last_price.get((item, vendor))
        ws.cell(row=row, column=sc, value=item)
        ws.cell(row=row, column=sc + 1, value=vendor or "-")
        if previous is not None:
            change = current - previous
            pct = (change / previous * 100) if previous else None
            direction = "↑" if change > 0 else ("↓" if change < 0 else "–")
            c_prev = ws.cell(row=row, column=sc + 2, value=previous)
            c_prev.number_format = MONEY_FORMAT
            c_cur = ws.cell(row=row, column=sc + 3, value=current)
            c_cur.number_format = MONEY_FORMAT
            c_chg = ws.cell(row=row, column=sc + 4, value=change)
            c_chg.number_format = MONEY_FORMAT
            if pct is not None:
                ws.cell(row=row, column=sc + 5, value=round(pct, 1))
            ws.cell(row=row, column=sc + 6, value=direction)
        else:
            c_cur = ws.cell(row=row, column=sc + 3, value=current)
            c_cur.number_format = MONEY_FORMAT
            ws.cell(row=row, column=sc + 6, value="new")
        row += 1

    for col in range(sc, sc + 7):
        ws.column_dimensions[get_column_letter(col)].width = 18

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


STOCK_MOVEMENT_HEADERS = ["Date", "Sender", "Item", "Quantity", "Unit"]


def _write_movement_sheet(ws, rows: list[dict]):
    for col, header in enumerate(STOCK_MOVEMENT_HEADERS, start=1):
        ws.cell(row=1, column=col, value=header).font = BOLD
    for r, row in enumerate(rows, start=2):
        ws.cell(row=r, column=1, value=row.get("created_at"))
        ws.cell(row=r, column=2, value=row.get("sender_phone"))
        ws.cell(row=r, column=3, value=row.get("item"))
        ws.cell(row=r, column=4, value=row.get("quantity"))
        ws.cell(row=r, column=5, value=row.get("unit"))
    for col in range(1, len(STOCK_MOVEMENT_HEADERS) + 1):
        ws.column_dimensions[get_column_letter(col)].width = 18


async def generate_stock_report_workbook(start: date, end: date) -> bytes:
    all_time = await get_all_stock_transactions()
    period_rows = await get_stock_transactions_between(start, end)

    balance = defaultdict(lambda: {"in": 0.0, "out": 0.0, "unit": None})
    for row in all_time:
        key = (row.get("item") or "").strip().lower()
        bucket = balance[key]
        bucket["unit"] = bucket["unit"] or row.get("unit")
        bucket[row["direction"]] += float(row["quantity"])

    wb = Workbook()
    ws_stock = wb.active
    ws_stock.title = "Current Stock"
    headers = ["Item", "Unit", "Total In", "Total Out", "Balance"]
    for col, header in enumerate(headers, start=1):
        ws_stock.cell(row=1, column=col, value=header).font = BOLD
    for r, (item, b) in enumerate(sorted(balance.items()), start=2):
        ws_stock.cell(row=r, column=1, value=item)
        ws_stock.cell(row=r, column=2, value=b["unit"])
        ws_stock.cell(row=r, column=3, value=b["in"])
        ws_stock.cell(row=r, column=4, value=b["out"])
        ws_stock.cell(row=r, column=5, value=b["in"] - b["out"])
    for col in range(1, len(headers) + 1):
        ws_stock.column_dimensions[get_column_letter(col)].width = 16

    in_rows = [r for r in period_rows if r["direction"] == "in"]
    out_rows = [r for r in period_rows if r["direction"] == "out"]

    ws_in = wb.create_sheet("IN")
    _write_movement_sheet(ws_in, in_rows)

    ws_out = wb.create_sheet("OUT")
    _write_movement_sheet(ws_out, out_rows)

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
